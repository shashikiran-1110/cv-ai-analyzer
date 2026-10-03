"""Tailored resume → .docx (ROADMAP §9.5). Keeps the resume's own line order; headings and bullets are styled."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt

from . import resume

_BULLET = re.compile(r"^\s*[-•*▪◦●·]\s+")


def resume_docx(text: str, title: str = "") -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    doc.core_properties.title = title[:200]
    lines = [l.rstrip() for l in text.splitlines()]
    first = True
    for line in lines:
        t = line.strip()
        if not t:
            continue
        if first:
            doc.add_heading(t[:120], level=0)
            first = False
            continue
        words = len(t.split())
        if words <= 4 and (resume._EXP_HEAD.match(t.strip(":")) or resume._OTHER_HEAD.match(t.strip(":"))):
            doc.add_heading(t.strip(":"), level=1)
        elif _BULLET.match(t):
            doc.add_paragraph(_BULLET.sub("", t), style="List Bullet")
        else:
            p = doc.add_paragraph()
            run = p.add_run(t)
            run.bold = bool(resume._RANGE.search(t)) and words <= 14     # role header lines with dates
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
